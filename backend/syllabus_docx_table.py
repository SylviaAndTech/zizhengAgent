"""
Word(.docx)课程大纲里"表5-1 课程知识单元、教学内容、教学目标及学时分配对应关系"和
"表5-2 实验单元、教学内容、教学目标、课程目标、类型、要求及学时分配对应关系"这两张表的解析。

跟PDF那两条路径（syllabus_table_ocr按坐标分列 / syllabus_vision_ocr视觉大模型）不同，
docx本身就有结构化的表格，直接用python-docx读单元格即可，不需要OCR、也不需要大模型。
原来的extract_text_from_docx只读document.paragraphs，表格里的文字根本读不到。

要处理的几个Word表格特有的问题（都是在真实大纲里遇到过的）：
1. "知识单元"列是纵向合并单元格，python-docx按行读时，每一行都会重复返回这个单元名；
   而"教学内容"列反而常被拆成好几行（编号条目一行、"重点："一行、"难点："一行）。
   所以按单元名对连续行分组，同组的教学内容按顺序拼接——一个知识单元/实验单元只对应
   知识点库里的一条记录（跟扫描版PDF路径的合并规则一致）。
2. 编号列表很多是Word自动编号（段落的numPr），para.text里读不到"1." "2."这些编号，
   要按numbering.xml里的编号格式自己补回来；同一格里自动编号和手打编号混用也很常见。
3. 表头文字带换行（"知识\\n单元"），匹配关键词前要去掉空白。

只处理这两张表，不是通用的Word表格解析器。
"""
import re

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

# "章节"列的表头关键词：理论课是表5-1"知识单元"+表5-2"实验单元"，综合实践课的表5-1是
# "实践单元"/"实验（实践）单元"，专业实训是"实训内容"，课程设计是"教学环节"。直接按表头
# 识别目标表，不依赖"表5-1/表5-2"这种标题编号——编号在不同课程大纲里会变，表头才是稳定的
_UNIT_HEADER_KEYWORDS = ("知识单元", "实验单元", "实践单元", "实训内容", "教学环节")
_CONTENT_HEADER_KEYWORD = "教学内容"
_TOTAL_ROW_RE = re.compile(r"^(合计|总计|小计)")
_CN_DIGITS = "零一二三四五六七八九"


def _norm(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _header_key(text: str) -> str:
    """表头匹配用：去空白、去括号注释——"实验\\n（实践）\\n单元"要能匹配上"实验单元"，
    "教学内容\\n（知识点）"匹配"教学内容\""""
    return re.sub(r"[（(][^）)]*[）)]", "", _norm(text))


def _header_columns(table: Table) -> tuple[int, int, int] | None:
    """在表格前两行里找"单元"列和"教学内容"列的列号，返回(表头行号, 单元列, 内容列)，
    不是目标表就返回None。不写死列号——表5-2比表5-1多了"实验类型/实验要求"等列"""
    for ri, row in enumerate(table.rows[:2]):
        texts = [_header_key(c.text) for c in row.cells]
        unit_col = next((i for i, t in enumerate(texts) if any(k in t for k in _UNIT_HEADER_KEYWORDS)), None)
        content_col = next((i for i, t in enumerate(texts) if _CONTENT_HEADER_KEYWORD in t), None)
        if unit_col is not None and content_col is not None and unit_col != content_col:
            return ri, unit_col, content_col
    return None


class _Numbering:
    """把Word自动编号还原成文字（"1." "2." ...）。
    计数器按(numId, ilvl)记，遇到上一级编号时下级计数器归零，跟Word的显示规则一致；
    只覆盖大纲里实际会用到的几种编号格式，其余格式退化成阿拉伯数字。"""

    def __init__(self, document):
        self._levels = {}  # (numId, ilvl) -> (numFmt, lvlText, start)
        self._counters = {}  # numId -> {ilvl: 当前值}
        try:
            numbering = document.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return

        abstract = {}
        for an in numbering.findall(qn("w:abstractNum")):
            lvls = {}
            for lvl in an.findall(qn("w:lvl")):
                ilvl = int(lvl.get(qn("w:ilvl")))
                lvls[ilvl] = (
                    self._val(lvl, "w:numFmt", "decimal"),
                    self._val(lvl, "w:lvlText", "%1."),
                    int(self._val(lvl, "w:start", "1")),
                )
            abstract[an.get(qn("w:abstractNumId"))] = lvls

        for num in numbering.findall(qn("w:num")):
            num_id = num.get(qn("w:numId"))
            abs_ref = num.find(qn("w:abstractNumId"))
            if abs_ref is None:
                continue
            lvls = dict(abstract.get(abs_ref.get(qn("w:val")), {}))
            for override in num.findall(qn("w:lvlOverride")):
                ilvl = int(override.get(qn("w:ilvl")))
                start = override.find(qn("w:startOverride"))
                if start is not None and ilvl in lvls:
                    fmt, text, _ = lvls[ilvl]
                    lvls[ilvl] = (fmt, text, int(start.get(qn("w:val"))))
            for ilvl, spec in lvls.items():
                self._levels[(num_id, ilvl)] = spec

    @staticmethod
    def _val(el, tag, default):
        child = el.find(qn(tag))
        return child.get(qn("w:val")) if child is not None and child.get(qn("w:val")) is not None else default

    @staticmethod
    def _num_pr(paragraph: Paragraph):
        """段落自身的numPr，没有就看段落样式上的numPr（列表样式的编号挂在样式上）"""
        ppr = paragraph._p.pPr
        if ppr is not None and ppr.numPr is not None:
            return ppr.numPr
        style = paragraph.style
        while style is not None:
            sppr = style.element.pPr
            if sppr is not None and sppr.numPr is not None:
                return sppr.numPr
            style = style.base_style
        return None

    @staticmethod
    def _format(value: int, fmt: str) -> str:
        if fmt in ("lowerLetter", "upperLetter"):
            s = chr(ord("a") + (value - 1) % 26)
            return s.upper() if fmt == "upperLetter" else s
        if fmt in ("chineseCounting", "chineseCountingThousand", "ideographTraditional") and 0 < value < 100:
            if value < 10:
                return _CN_DIGITS[value]
            tens, ones = divmod(value, 10)
            return ("" if tens == 1 else _CN_DIGITS[tens]) + "十" + (_CN_DIGITS[ones] if ones else "")
        if fmt == "decimalEnclosedCircle" and 0 < value <= 20:
            return chr(0x2460 + value - 1)
        return str(value)

    def prefix(self, paragraph: Paragraph) -> str:
        num_pr = self._num_pr(paragraph)
        if num_pr is None or num_pr.numId is None:
            return ""
        num_id = str(num_pr.numId.val)
        if num_id == "0":  # numId=0表示"去掉编号"
            return ""
        ilvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        spec = self._levels.get((num_id, ilvl))
        if spec is None:
            return ""
        fmt, lvl_text, start = spec
        if fmt in ("none", "bullet"):
            return ""

        counters = self._counters.setdefault(num_id, {})
        counters[ilvl] = counters.get(ilvl, start - 1) + 1
        for deeper in [k for k in counters if k > ilvl]:
            del counters[deeper]

        def repl(m):
            level = int(m.group(1)) - 1
            level_spec = self._levels.get((num_id, level))
            value = counters.get(level, level_spec[2] if level_spec else 1)
            return self._format(value, level_spec[0] if level_spec else "decimal")

        return re.sub(r"%(\d)", repl, lvl_text)


def _cell_text(cell, numbering: _Numbering) -> str:
    lines = []
    for para in cell.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        lines.append(numbering.prefix(para) + text)
    return "\n".join(lines)


def extract_units_and_points(file_path: str) -> list[tuple[str, str]]:
    """
    返回 [(章节, 知识点描述), ...]：表5-1在前、表5-2在后，按表格里出现的顺序。
    章节=知识单元/实验单元/实践单元列的原文，知识点描述=该单元对应的"教学内容"整格原文
    （编号、重点、难点都保留，补回自动编号），一个单元一条记录。
    一张目标表都没找到时返回空列表，调用方自行回退到其他解析方式。
    """
    document = docx.Document(file_path)
    numbering = _Numbering(document)
    groups: list[list] = []  # [[单元名, [教学内容片段, ...]], ...]
    seen_tcs = set()

    # document.tables按文档顺序排列，表5-1自然在表5-2前面；跨页被拆成两段的续表表头
    # 重复出现，也会被识别成目标表，分组状态跨表格保留，被拆开的同一个单元接着拼
    for table in document.tables:
        header = _header_columns(table)
        if header is None:
            continue
        header_row, unit_col, content_col = header

        for row in table.rows[header_row + 1:]:
            cells = row.cells
            if len(cells) <= max(unit_col, content_col):
                continue
            unit = re.sub(r"\s*\n\s*", "", cells[unit_col].text.strip())
            if not unit or _TOTAL_ROW_RE.match(unit):
                continue
            if not groups or groups[-1][0] != unit:
                groups.append([unit, []])

            content_cell = cells[content_col]
            # 教学内容列如果也是纵向合并的，每一行会重复返回同一个格子，只取一次；
            # 同时也避免同一个自动编号段落被计数两次
            if content_cell._tc in seen_tcs:
                continue
            seen_tcs.add(content_cell._tc)
            text = _cell_text(content_cell, numbering)
            if text:
                groups[-1][1].append(text)

    return [(unit, "\n".join(parts)) for unit, parts in groups if parts]
